"""A deliberately narrow DOCX profile, with independent actual-file XML readback.

SDT tags carry block IDs only; text, structure, list markers, cell values, and link
targets are always reconstructed from document.xml and its relationships. This is
not a general Word importer: unrecognized content is an error, never skipped.
"""
from __future__ import annotations

import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile, ZipInfo, ZIP_DEFLATED

from .errors import AssemblyError
from .paths import checked_path, read_bytes

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
XML = "http://www.w3.org/XML/1998/namespace"
TAG_PREFIX = "final-assembly:"
STYLES = {"Normal", "Title", "FAListBullet", "FAListNumber", "FAQuote", "FACode", *[f"Heading{i}" for i in range(1, 7)]}
PACKAGE_PARTS = {
    "[Content_Types].xml", "_rels/.rels", "docProps/core.xml", "docProps/app.xml",
    "docProps/thumbnail.jpeg", "word/document.xml", "word/_rels/document.xml.rels",
    "word/styles.xml", "word/settings.xml", "word/webSettings.xml",
    "word/fontTable.xml", "word/theme/theme1.xml",
}


def _w(name):
    return f"{{{W}}}{name}"


def _fail(message, location, code="DOCX_UNSUPPORTED"):
    raise AssemblyError(code, message, location)


def _valid_url(url, location):
    if not isinstance(url, str) or any(ord(c) < 33 or ord(c) == 127 for c in url):
        _fail("Invalid hyperlink target", location, "DOCX_LINK_INVALID")
    try:
        parsed = urlsplit(url)
        valid = (parsed.scheme in {"http", "https"} and bool(parsed.hostname)) or (parsed.scheme == "mailto" and bool(parsed.path))
    except ValueError:
        valid = False
    if not valid:
        _fail("Only absolute http, https, and mailto links are supported", location, "DOCX_LINK_INVALID")
    return url


def write_docx(blocks: list[dict], path: Path) -> None:
    """Create readable DOCX; callers must verify the saved file before success."""
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise AssemblyError("DOCX_DEPENDENCY_MISSING", "Install python-docx from requirements.txt to build the verified Markdown/DOCX delivery") from exc

    path = checked_path(path)
    if path.exists():
        raise AssemblyError("OUTPUT_EXISTS", "Refusing to overwrite an existing output", str(path))
    document = Document()
    document.core_properties.title = ""
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    document.core_properties.subject = ""
    # Content/manifest reports carry real build provenance. Fixed package dates
    # eliminate incidental ZIP and template timestamps from reproducible output.
    document.core_properties.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.core_properties.revision = 1
    for key, rel in list(document.part.rels.items()):
        if rel.reltype.rsplit("/", 1)[-1] in {"customXml", "numbering", "stylesWithEffects"}:
            del document.part.rels[key]
    # Rebuild only used styles. This removes template title rules, theme colors,
    # hidden numbering, and unused styles that a verifier would need to interpret.
    for element in list(document.styles.element):
        if element.tag != qn("w:style") or element.get(qn("w:styleId")) not in STYLES:
            document.styles.element.remove(element)
    for name in ("FAListBullet", "FAListNumber", "FAQuote", "FACode"):
        document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    for style in document.styles:
        for child in list(style.element):
            if child.tag != qn("w:name"):
                style.element.remove(child)
        if style.style_id != "Normal":
            style.base_style = document.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.18
        if style.style_id == "FAQuote":
            style.paragraph_format.left_indent = Inches(.25)
            style.font.color.rgb = RGBColor.from_string("444444")
        if style.style_id == "FACode":
            style.font.name = "Consolas"
            style.font.size = Pt(9)
            style.paragraph_format.line_spacing = 1.05
            style.paragraph_format.space_before = Pt(6)
        if style.style_id in {"FAListBullet", "FAListNumber"}:
            style.paragraph_format.left_indent = Inches(.23)
            style.paragraph_format.first_line_indent = Inches(-.18)
            style.paragraph_format.space_after = Pt(5)
        if style.style_id == "Title" or style.style_id.startswith("Heading"):
            level = 1 if style.style_id == "Title" else int(style.style_id[-1])
            style.font.size = Pt(19 if style.style_id == "Title" else max(11, 17 - level))
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(0 if style.style_id == "Title" else 11)
            style.paragraph_format.space_after = Pt(7)
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), str(level - 1))
            style.element.get_or_add_pPr().append(outline)
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.7), Inches(.7)
    section.left_margin, section.right_margin = Inches(.7), Inches(.7)

    def add_spans(paragraph, spans):
        for span in spans:
            if not span["text"]:
                continue
            run = paragraph.add_run(span["text"])
            for key, attribute in (("bold", "bold"), ("italic", "italic"), ("strike", "strike")):
                if span.get(key):
                    setattr(run.font, attribute, True)
            if span.get("code"):
                run.font.name = "Consolas"
            if "url" not in span:
                continue
            url = _valid_url(span["url"], "export.link")
            relationship = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), relationship)
            run.font.color.rgb = RGBColor.from_string("174B73")
            run.font.underline = True
            link.append(run._r)
            paragraph._p.append(link)

    seen = set()
    first_unit = True
    for block in blocks:
        block_id = block["id"]
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,79}", block_id) or block_id in seen:
            _fail("Invalid or duplicate block ID", block_id, "DOCX_BLOCK_INVALID")
        seen.add(block_id)
        sdt, properties, content = OxmlElement("w:sdt"), OxmlElement("w:sdtPr"), OxmlElement("w:sdtContent")
        tag = OxmlElement("w:tag")
        tag.set(qn("w:val"), TAG_PREFIX + block_id)
        properties.append(tag)
        sdt.extend([properties, content])
        document.element.body.insert(len(document.element.body) - 1, sdt)
        for unit in block["units"]:
            kind = unit["kind"]
            if kind == "table":
                rows = unit["rows"]
                if not rows or not rows[0] or any(len(row) != len(rows[0]) for row in rows):
                    _fail("Tables must be nonempty and rectangular", block_id, "DOCX_TABLE_INVALID")
                table = document.add_table(rows=len(rows), cols=len(rows[0]))
                table.autofit = False
                # Widths reflect the actual CJK/Latin lengths instead of equal
                # widths; cap extreme lengths so a long cell cannot crush peers.
                weights = [max(4, min(28, max(sum(2 if ord(ch) > 255 else 1 for span in row[col] for ch in span["text"]) for row in rows))) for col in range(len(rows[0]))]
                widths = [Inches(7.1 * weight / sum(weights)) for weight in weights]
                for column, width in zip(table.columns, widths):
                    column.width = width
                props = table._tbl.tblPr
                borders = OxmlElement("w:tblBorders")
                for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                    border = OxmlElement(f"w:{edge}")
                    for attr, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                        border.set(qn(f"w:{attr}"), val)
                    borders.append(border)
                props.append(borders)
                margins = OxmlElement("w:tblCellMar")
                for edge, val in (("top", "90"), ("bottom", "90"), ("left", "105"), ("right", "105")):
                    margin = OxmlElement(f"w:{edge}")
                    margin.set(qn("w:w"), val)
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                props.append(margins)
                for row_index, (row, docx_row) in enumerate(zip(rows, table.rows, strict=True)):
                    if row_index == 0:
                        docx_row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                    # Table.cell() reconstructs the complete table grid on every
                    # access. This writer creates rectangular, unmerged rows, so
                    # materialize each row's cell proxies once and reuse them.
                    row_cells = tuple(docx_row.cells)
                    for col_index, spans in enumerate(row):
                        cell = row_cells[col_index]
                        cell.width = widths[col_index]
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        paragraph = cell.paragraphs[0]
                        paragraph.style = document.styles["Normal"]
                        paragraph.paragraph_format.space_after = Pt(1)
                        paragraph.paragraph_format.space_before = Pt(1)
                        paragraph.paragraph_format.line_spacing = 1.15
                        align = unit.get("alignment", ["left"] * len(row))[col_index]
                        paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
                        add_spans(paragraph, spans)
                        if row_index == 0:
                            shade = OxmlElement("w:shd")
                            shade.set(qn("w:fill"), "DFEAF3")
                            cell._tc.get_or_add_tcPr().append(shade)
                        for run in paragraph.runs:
                            run.font.size = Pt(10.5)
                content.append(table._tbl)
            elif kind in {"paragraph", "heading", "list_item", "quote", "code_block"}:
                style = "Normal"
                if kind == "heading":
                    level = unit["level"]
                    if not isinstance(level, int) or not 1 <= level <= 6:
                        _fail("Heading level must be 1 through 6", block_id)
                    style = "Title" if first_unit and level == 1 else f"Heading {level}"
                elif kind == "list_item":
                    style = "FAListNumber" if unit["ordered"] else "FAListBullet"
                elif kind in {"quote", "code_block"}:
                    style = "FAQuote" if kind == "quote" else "FACode"
                paragraph = document.add_paragraph(style=style)
                if kind == "list_item":
                    paragraph.add_run(f"{unit['number']}. " if unit["ordered"] else "• ")
                add_spans(paragraph, unit["spans"])
                content.append(paragraph._p)
            else:
                _fail(f"Unsupported unit kind: {kind}", block_id)
            first_unit = False
    path.parent.mkdir(parents=True, exist_ok=True)
    buffer = BytesIO()
    document.save(buffer)
    # Exclusive open closes the race between the exists check and the save.
    with ZipFile(BytesIO(buffer.getvalue())) as source, path.open("xb") as stream:
        with ZipFile(stream, "w", compression=ZIP_DEFLATED, compresslevel=9) as output:
            for name in sorted(source.namelist()):
                entry = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
                entry.compress_type = ZIP_DEFLATED
                entry.create_system = 3
                entry.external_attr = 0o600 << 16
                output.writestr(entry, source.read(name), compresslevel=9)


def _xml(data, location):
    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        _fail("DTD/entity declarations are not accepted", location, "DOCX_XML_INVALID")
    try:
        return ET.fromstring(data)
    except ET.ParseError as exc:
        _fail(f"Malformed XML: {exc}", location, "DOCX_XML_INVALID")


def _children(node, allowed, location, repeats=()):
    """Allowlist each direct child, including properties, without hiding text."""
    counts = {}
    for child in node:
        if child.tag not in {_w(name) for name in allowed}:
            _fail(f"Unsupported XML element {child.tag}", location)
        name = child.tag.rsplit("}", 1)[-1]
        counts[name] = counts.get(name, 0) + 1
        if counts[name] > 1 and name not in repeats:
            _fail(f"Duplicate XML element {name}", location, "DOCX_XML_INVALID")
    if node.text and node.text.strip():
        _fail("Unexpected XML text outside a text run", location)
    if any(child.tail and child.tail.strip() for child in node):
        _fail("Unexpected text after an XML element", location)


def _leaf_properties(node, names, location):
    if node is None:
        return
    _children(node, names, location)
    for child in node:
        if len(child) or (child.text and child.text.strip()):
            _fail("Unsupported nested property content", location)


def _rpr(node, location):
    _leaf_properties(node, {"rFonts", "b", "bCs", "i", "strike", "color", "sz", "szCs", "u", "lang"}, location)


def _ppr(node, location):
    _leaf_properties(node, {"pStyle", "keepNext", "keepLines", "spacing", "ind", "jc", "outlineLvl", "widowControl"}, location)


def _append_span(spans, text, url=None, formatting=None):
    if not text:
        return
    attributes = {**(formatting or {}), **({"url": url} if url is not None else {})}
    if spans and {key: value for key, value in spans[-1].items() if key != "text"} == attributes:
        spans[-1]["text"] += text
    else:
        spans.append({"text": text, **attributes})


def _run_format(run):
    properties = run.find(_w("rPr"))
    result = {}
    if properties is not None:
        for name, key in (("b", "bold"), ("i", "italic"), ("strike", "strike")):
            node = properties.find(_w(name))
            if node is not None and node.get(_w("val"), "true") in {"1", "true", "on"}:
                result[key] = True
        fonts = properties.find(_w("rFonts"))
        if fonts is not None and fonts.get(_w("ascii")) == "Consolas" and fonts.get(_w("hAnsi")) == "Consolas":
            result["code"] = True
    return result


def _read_run(run, location):
    _children(run, {"rPr", "t", "tab", "br"}, location, repeats={"t", "tab", "br"})
    _rpr(run.find(_w("rPr")), location)
    texts = []
    for child in run:
        if child.tag == _w("rPr"):
            continue
        if len(child):
            _fail("Nested run content is unsupported", location)
        if child.tag == _w("t"):
            if set(child.attrib) - {f"{{{XML}}}space"}:
                _fail("Unsupported text-run attributes", location)
            text = child.text or ""
            whitespace_mode = child.get(f"{{{XML}}}space", "default")
            if whitespace_mode not in {"default", "preserve"} or (text != text.strip() and whitespace_mode != "preserve"):
                _fail("Significant edge whitespace requires xml:space=preserve", location, "DOCX_XML_INVALID")
            texts.append(text)
        elif child.tag == _w("tab"):
            if child.attrib or child.text:
                _fail("Unsupported tab content", location)
            texts.append("\t")
        else:
            if set(child.attrib) - {_w("type")} or child.get(_w("type"), "textWrapping") != "textWrapping" or child.text:
                _fail("Only text-wrapping line breaks are supported", location)
            texts.append("\n")
    return "".join(texts)


def _read_paragraph(paragraph, relationships, used_links, location, cell=False):
    _children(paragraph, {"pPr", "r", "hyperlink"}, location, repeats={"r", "hyperlink"})
    properties = paragraph.find(_w("pPr"))
    _ppr(properties, location)
    style_node = properties.find(_w("pStyle")) if properties is not None else None
    style = "Normal" if style_node is None else style_node.get(_w("val"))
    if style not in STYLES:
        _fail(f"Unknown paragraph style: {style}", location, "DOCX_STYLE_INVALID")
    if properties is not None and properties.find(_w("outlineLvl")) is not None:
        _fail("Direct outline-level overrides are unsupported", location, "DOCX_STYLE_INVALID")
    spans = []
    for child in paragraph:
        if child.tag == _w("r"):
            _append_span(spans, _read_run(child, location), formatting=_run_format(child))
        elif child.tag == _w("hyperlink"):
            if set(child.attrib) - {f"{{{R}}}id", _w("history")}:
                _fail("Only external relationship hyperlinks are supported", location, "DOCX_LINK_INVALID")
            relation_id = child.get(f"{{{R}}}id")
            if relation_id not in relationships:
                _fail(f"Missing hyperlink relationship: {relation_id}", location, "DOCX_LINK_INVALID")
            used_links.add(relation_id)
            _children(child, {"r"}, location, repeats={"r"})
            if not len(child):
                _fail("Empty hyperlink label", location, "DOCX_LINK_INVALID")
            for run in child:
                _append_span(spans, _read_run(run, location), relationships[relation_id], _run_format(run))
    if cell:
        if style != "Normal":
            _fail("Table cells must contain a plain paragraph", location, "DOCX_STYLE_INVALID")
        return spans
    if style in {"FAListBullet", "FAListNumber"}:
        if not spans or "url" in spans[0]:
            _fail("List marker is missing", location, "DOCX_LIST_INVALID")
        ordered = style == "FAListNumber"
        match = re.match(r"([0-9]+)\. " if ordered else r"• ", spans[0]["text"])
        if not match:
            _fail("Actual list marker disagrees with its paragraph style", location, "DOCX_LIST_INVALID")
        marker = match.group(0)
        spans[0]["text"] = spans[0]["text"][len(marker):]
        if not spans[0]["text"]:
            spans.pop(0)
        return {"kind": "list_item", "ordered": ordered, **({"number": int(match.group(1))} if ordered else {}), "spans": spans}
    if style == "Title" or style.startswith("Heading"):
        return {"kind": "heading", "level": 1 if style == "Title" else int(style[-1]), "spans": spans}
    if style in {"FAQuote", "FACode"}:
        return {"kind": "quote" if style == "FAQuote" else "code_block", "spans": spans}
    return {"kind": "paragraph", "spans": spans}


def _read_table(table, relationships, used_links, location):
    _children(table, {"tblPr", "tblGrid", "tr"}, location, repeats={"tr"})
    props = table.find(_w("tblPr"))
    if props is None:
        _fail("Missing table properties", location)
    _children(props, {"tblW", "tblLayout", "tblLook", "tblBorders", "tblCellMar"}, location)
    for prop in props:
        name = prop.tag.rsplit("}", 1)[-1]
        if name == "tblBorders":
            _leaf_properties(prop, {"top", "left", "bottom", "right", "insideH", "insideV"}, location)
        elif name == "tblCellMar":
            _leaf_properties(prop, {"top", "left", "bottom", "right"}, location)
        elif len(prop) or prop.text:
            _fail("Unsupported table property content", location)
    grid = table.find(_w("tblGrid"))
    if grid is None or not len(grid):
        _fail("Table grid is missing", location, "DOCX_TABLE_INVALID")
    _children(grid, {"gridCol"}, location, repeats={"gridCol"})
    if any(len(col) or col.text for col in grid):
        _fail("Invalid table grid", location)
    rows, alignment = [], []
    for row_index, row in enumerate(table.findall(_w("tr")), 1):
        row_location = f"{location}.rows[{row_index}]"
        _children(row, {"trPr", "tc"}, row_location, repeats={"tc"})
        row_cells = row.findall(_w("tc"))
        if len(row_cells) != len(grid):
            _fail(f"Table row has {len(row_cells)} cells; expected {len(grid)}", row_location, "DOCX_TABLE_INVALID")
        row_props = row.find(_w("trPr"))
        _leaf_properties(row_props, {"tblHeader", "cantSplit"}, row_location)
        header = None if row_props is None else row_props.find(_w("tblHeader"))
        is_header = header is not None and header.get(_w("val"), "true") in {"1", "true", "on"}
        if is_header != (row_index == 1):
            _fail("Only the first table row must be marked as its header", row_location, "DOCX_TABLE_INVALID")
        cells = []
        for column, cell in enumerate(row_cells, 1):
            cell_location = f"{row_location}.cells[{column}]"
            _children(cell, {"tcPr", "p"}, cell_location)
            _leaf_properties(cell.find(_w("tcPr")), {"tcW", "vAlign", "shd"}, cell_location)
            paragraph = cell.find(_w("p"))
            if paragraph is None:
                _fail("A table cell needs exactly one paragraph", cell_location, "DOCX_TABLE_INVALID")
            cells.append(_read_paragraph(paragraph, relationships, used_links, cell_location, cell=True))
            align_node = paragraph.find(f"{_w('pPr')}/{_w('jc')}")
            align = align_node.get(_w("val"), "left") if align_node is not None else "left"
            if align not in {"left", "center", "right"}:
                _fail("Unsupported table alignment", cell_location, "DOCX_TABLE_INVALID")
            if row_index == 1:
                alignment.append(align)
            elif align != alignment[column - 1]:
                _fail("Table column alignment changed between rows", cell_location, "DOCX_TABLE_INVALID")
        rows.append(cells)
    if not rows:
        _fail("Table has no rows", location, "DOCX_TABLE_INVALID")
    return {"kind": "table", "rows": rows, "alignment": alignment}


def _check_styles(styles):
    if styles.tag != _w("styles"):
        _fail("Invalid styles root", "word/styles.xml")
    _children(styles, {"style"}, "word/styles.xml", repeats={"style"})
    seen = set()
    for style in styles:
        style_id = style.get(_w("styleId"))
        location = f"styles.{style_id}"
        if style_id not in STYLES or style_id in seen or style.get(_w("type")) != "paragraph":
            _fail("Unknown, duplicate, or non-paragraph style", location, "DOCX_STYLE_INVALID")
        seen.add(style_id)
        _children(style, {"name", "basedOn", "pPr", "rPr"}, location)
        for property_name in ("name", "basedOn"):
            property_node = style.find(_w(property_name))
            if property_node is not None and (len(property_node) or (property_node.text and property_node.text.strip())):
                _fail("Unexpected content inside a style property", location)
        base = style.find(_w("basedOn"))
        if (style_id == "Normal" and base is not None) or (style_id != "Normal" and (base is None or base.get(_w("val")) != "Normal")):
            _fail("Unexpected style inheritance", location, "DOCX_STYLE_INVALID")
        _ppr(style.find(_w("pPr")), location)
        _rpr(style.find(_w("rPr")), location)
        semantic_style = _run_format(style)
        expected_style = {"bold": True} if style_id == "Title" or style_id.startswith("Heading") else {"code": True} if style_id == "FACode" else {}
        if semantic_style != expected_style:
            _fail("Default semantic formatting changed", location, "DOCX_STYLE_INVALID")
        if style.find(f"{_w('pPr')}/{_w('pStyle')}") is not None:
            _fail("A style must not redirect to another paragraph style", location, "DOCX_STYLE_INVALID")
        outline = style.find(f"{_w('pPr')}/{_w('outlineLvl')}")
        expected = "0" if style_id == "Title" else str(int(style_id[-1]) - 1) if style_id.startswith("Heading") else None
        actual = outline.get(_w("val")) if outline is not None else None
        if actual != expected:
            _fail("Style outline level changed", location, "DOCX_STYLE_INVALID")
    if seen != STYLES:
        _fail("Required profile styles are missing", "word/styles.xml", "DOCX_STYLE_INVALID")


def _check_package_metadata(root_rels, content_types, names):
    """Confirm the ZIP's actual main document is the one we are inspecting."""
    location = "package.relationships"
    if root_rels.tag != f"{{{PKG}}}Relationships":
        _fail("Invalid package relationships root", location, "DOCX_PACKAGE_INVALID")
    allowed = {
        f"{R}/officeDocument": "word/document.xml",
        f"{PKG}/metadata/core-properties": "docProps/core.xml",
        f"{R}/extended-properties": "docProps/app.xml",
        f"{PKG}/metadata/thumbnail": "docProps/thumbnail.jpeg",
    }
    seen_ids, seen_types = set(), set()
    for relation in root_rels:
        rid, kind, target = (relation.get(key) for key in ("Id", "Type", "Target"))
        if (relation.tag != f"{{{PKG}}}Relationship" or len(relation) or
                set(relation.attrib) != {"Id", "Type", "Target"} or not rid or
                rid in seen_ids or kind in seen_types or kind not in allowed or
                target != allowed[kind] or target not in names):
            _fail("Unsupported package relationship or main-document target", location, "DOCX_PACKAGE_INVALID")
        seen_ids.add(rid)
        seen_types.add(kind)
    if f"{R}/officeDocument" not in seen_types:
        _fail("Missing main-document relationship", location, "DOCX_PACKAGE_INVALID")
    namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    if content_types.tag != f"{{{namespace}}}Types":
        _fail("Invalid content-types root", "package.content_types", "DOCX_PACKAGE_INVALID")
    base = "application/vnd.openxmlformats-officedocument."
    overrides = {
        "/docProps/app.xml": base + "extended-properties+xml",
        "/docProps/core.xml": "application/vnd.openxmlformats-package.core-properties+xml",
        "/word/document.xml": base + "wordprocessingml.document.main+xml",
        "/word/styles.xml": base + "wordprocessingml.styles+xml",
        "/word/settings.xml": base + "wordprocessingml.settings+xml",
        "/word/webSettings.xml": base + "wordprocessingml.webSettings+xml",
        "/word/fontTable.xml": base + "wordprocessingml.fontTable+xml",
        "/word/theme/theme1.xml": base + "theme+xml",
    }
    defaults = {"jpeg": "image/jpeg", "rels": "application/vnd.openxmlformats-package.relationships+xml", "xml": "application/xml"}
    seen = set()
    for node in content_types:
        if node.tag == f"{{{namespace}}}Override":
            name, expected = node.get("PartName"), overrides
            attributes = {"PartName", "ContentType"}
            if name and name.lstrip("/") not in names:
                _fail("Content type refers to a missing package part", "package.content_types", "DOCX_PACKAGE_INVALID")
        elif node.tag == f"{{{namespace}}}Default":
            name, expected = node.get("Extension"), defaults
            attributes = {"Extension", "ContentType"}
        else:
            _fail("Unsupported content-type element", "package.content_types", "DOCX_PACKAGE_INVALID")
        if name in seen or name not in expected or node.get("ContentType") != expected[name] or set(node.attrib) != attributes or len(node):
            _fail("Unsupported, duplicate, or changed content type", "package.content_types", "DOCX_PACKAGE_INVALID")
        seen.add(name)
    if not {"/word/document.xml", "/word/styles.xml", "rels", "xml"} <= seen:
        _fail("Required package content types are missing", "package.content_types", "DOCX_PACKAGE_INVALID")


def read_docx_bytes(raw: bytes) -> list[dict]:
    """Read one immutable saved-file snapshot with zipfile + ElementTree.

    Word resaves that introduce unsupported OOXML are deliberately rejected.
    Content normalization does not replace separate renderer-based visual QA.
    """
    if not isinstance(raw, bytes) or len(raw) > 16 * 1024 * 1024:
        raise AssemblyError("DOCX_PACKAGE_INVALID", "Expected a DOCX byte snapshot of at most 16 MiB", "docx")
    try:
        with ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            names = [info.filename for info in entries]
            if len(names) != len(set(names)):
                _fail("Duplicate ZIP entry names", "docx", "DOCX_PACKAGE_INVALID")
            if set(names) - PACKAGE_PARTS:
                _fail(f"Unsupported package parts: {sorted(set(names) - PACKAGE_PARTS)}", "docx")
            if sum(info.file_size for info in entries) > 20 * 1024 * 1024 or any(info.flag_bits & 1 for info in entries):
                _fail("Oversized or encrypted DOCX package", "docx", "DOCX_PACKAGE_INVALID")
            document = _xml(archive.read("word/document.xml"), "word/document.xml")
            relationships_xml = _xml(archive.read("word/_rels/document.xml.rels"), "word/_rels/document.xml.rels")
            styles = _xml(archive.read("word/styles.xml"), "word/styles.xml")
            root_rels = _xml(archive.read("_rels/.rels"), "_rels/.rels")
            content_types = _xml(archive.read("[Content_Types].xml"), "[Content_Types].xml")
    except (BadZipFile, KeyError, OSError, RuntimeError) as exc:
        raise AssemblyError("DOCX_PACKAGE_INVALID", f"Cannot read DOCX: {exc}", "docx") from exc
    _check_package_metadata(root_rels, content_types, names)
    _check_styles(styles)
    relationships = {}
    relation_ids = set()
    expected_internal = {
        "styles": "styles.xml", "settings": "settings.xml", "webSettings": "webSettings.xml",
        "fontTable": "fontTable.xml", "theme": "theme/theme1.xml",
    }
    internal_seen = set()
    if relationships_xml.tag != f"{{{PKG}}}Relationships":
        _fail("Invalid relationship root", "relationships")
    for relation in relationships_xml:
        if relation.tag != f"{{{PKG}}}Relationship" or len(relation) or set(relation.attrib) - {"Id", "Type", "Target", "TargetMode"}:
            _fail("Unexpected relationship content", "relationships")
        rid, reltype, target = (relation.get(name) for name in ("Id", "Type", "Target"))
        if not rid or rid in relation_ids:
            _fail("Missing or duplicate relationship ID", "relationships", "DOCX_LINK_INVALID")
        relation_ids.add(rid)
        if reltype == f"{R}/hyperlink":
            if relation.get("TargetMode") != "External":
                _fail("Hyperlinks must be external relationships", f"relationships.{rid}", "DOCX_LINK_INVALID")
            relationships[rid] = _valid_url(target, f"relationships.{rid}")
        else:
            relname = reltype[len(R) + 1:] if reltype and reltype.startswith(R + "/") else None
            if relname not in expected_internal or target != expected_internal[relname] or relation.get("TargetMode") or relname in internal_seen or "word/" + target not in names:
                _fail("Unexpected internal or external relationship", f"relationships.{rid}")
            internal_seen.add(relname)
    if internal_seen != set(expected_internal):
        _fail("Required document relationships are missing", "relationships", "DOCX_PACKAGE_INVALID")
    if document.tag != _w("document"):
        _fail("Invalid document root", "document")
    _children(document, {"body"}, "document")
    body = document.find(_w("body"))
    if body is None:
        _fail("Missing document body", "document")
    _children(body, {"sdt", "sectPr"}, "document.body", repeats={"sdt"})
    section = body.find(_w("sectPr"))
    if section is None or list(body)[-1] is not section:
        _fail("The body must end with a single section definition", "document.body")
    _leaf_properties(section, {"pgSz", "pgMar", "cols", "docGrid"}, "document.section")
    blocks, block_ids, used_links = [], set(), set()
    for sdt in body.findall(_w("sdt")):
        _children(sdt, {"sdtPr", "sdtContent"}, "document.block")
        props, content = sdt.find(_w("sdtPr")), sdt.find(_w("sdtContent"))
        if props is None or content is None:
            _fail("Incomplete block wrapper", "document.block", "DOCX_BLOCK_INVALID")
        _leaf_properties(props, {"tag"}, "document.block")
        tag = props.find(_w("tag"))
        tag_value = "" if tag is None else tag.get(_w("val"), "")
        if not tag_value.startswith(TAG_PREFIX):
            _fail("Block tag is missing or unrecognized", "document.block", "DOCX_BLOCK_INVALID")
        block_id = tag_value[len(TAG_PREFIX):]
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,79}", block_id) or block_id in block_ids:
            _fail("Invalid or duplicate block ID", block_id, "DOCX_BLOCK_INVALID")
        block_ids.add(block_id)
        _children(content, {"p", "tbl"}, block_id, repeats={"p", "tbl"})
        units = []
        for unit_index, node in enumerate(content, 1):
            location = f"{block_id}.units[{unit_index}]"
            units.append(_read_paragraph(node, relationships, used_links, location) if node.tag == _w("p") else _read_table(node, relationships, used_links, location))
        blocks.append({"id": block_id, "units": units})
    return blocks


def read_docx(path: Path) -> list[dict]:
    """Path wrapper that reads the actual file exactly once under path policy."""
    return read_docx_bytes(read_bytes(path))
